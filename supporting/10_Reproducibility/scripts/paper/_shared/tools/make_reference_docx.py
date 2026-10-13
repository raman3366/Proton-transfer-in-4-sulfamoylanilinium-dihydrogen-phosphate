#!/usr/bin/env python3
"""Create the Word reference document (styles, page set-up, footer) that pandoc uses for every manuscript.

  tools/paper-env/bin/python paper/_shared/tools/make_reference_docx.py [--spacing 1.5] [--out PATH]

Starts from pandoc's own default reference.docx (so every style pandoc emits exists) and edits it with
python-docx:
  * Times New Roman 12 pt everywhere (theme fonts removed), en-GB proofing language
  * line spacing --spacing (1.5 by default; 2 for submission), left-aligned text, 6 pt after paragraphs
  * bold black headings (12 pt; level 3 bold italic), title 16 pt bold; no automatic heading numbers
  * captions 10 pt upright; "Table Text" 10 pt single-spaced for table cells; "Equation"/"Equation Number"
  * figures centred; table style with only top, header and bottom rules
  * A4, 25 mm margins, centred page number in the footer, continuous line numbering
Default output: paper/_shared/reference.docx (copied unchanged from cqc-dft).
"""
import argparse
import copy
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[2]
PANDOC = REPO / 'tools' / 'bin' / 'pandoc'
DEFAULT_OUT = REPO / 'paper' / '_shared' / 'reference.docx'
FONT = 'Times New Roman'


def set_fonts(rpr_parent, size_pt=None):
    """Explicit Times New Roman on an rPr (theme attributes removed, since they override w:ascii)."""
    rpr = rpr_parent.get_or_add_rPr() if hasattr(rpr_parent, 'get_or_add_rPr') else rpr_parent
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    for att in list(rfonts.attrib):
        if att.endswith('Theme') or att.endswith('theme'):
            del rfonts.attrib[att]
    for att in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(att), FONT)
    if size_pt is not None:
        for tag in ('w:sz', 'w:szCs'):
            el = rpr.find(qn(tag))
            if el is None:
                el = OxmlElement(tag)
                rpr.append(el)
            el.set(qn('w:val'), str(int(round(size_pt * 2))))


def strip_colour(style):
    rpr = style.element.find(qn('w:rPr'))
    if rpr is not None:
        for c in rpr.findall(qn('w:color')):
            rpr.remove(c)


def find_style(doc, name):
    # by UI name, then by style id (pandoc's 'Caption' has no lowercase built-in alias, so python-docx's
    # own lookup would warn)
    for st in doc.styles:
        if st.name == name:
            return st
    sid = name.replace(' ', '')
    for st in doc.styles:
        if st.style_id == sid:
            return st
    raise KeyError(name)


def get_style(doc, name, kind=WD_STYLE_TYPE.PARAGRAPH, base=None):
    try:
        return find_style(doc, name)
    except KeyError:
        st = doc.styles.add_style(name, kind)
        if base:
            st.base_style = find_style(doc, base)
        st.quick_style = True
        return st


def para_style(doc, name, size=None, bold=None, italic=None, align=None, before=None, after=None,
               line=None, keep_next=None, base=None):
    st = get_style(doc, name, base=base)
    set_fonts(st.element, size)
    f = st.font
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    f.color.rgb = None
    strip_colour(st)
    pf = st.paragraph_format
    if align is not None:
        pf.alignment = align
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
        pf.line_spacing = line
    if keep_next is not None:
        pf.keep_with_next = keep_next
    return st


def border(tag, val='single', sz=8):
    el = OxmlElement(tag)
    el.set(qn('w:val'), val)
    if val != 'nil':
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), 'auto')
    return el


def three_rule_table_style(doc):
    """Table style 'Table' (the one pandoc assigns): top and bottom rules, rule under the header row."""
    st = doc.styles['Table']
    el = st.element
    tblpr = el.find(qn('w:tblPr'))
    for old in tblpr.findall(qn('w:tblBorders')):
        tblpr.remove(old)
    b = OxmlElement('w:tblBorders')
    for tag, val, sz in (('w:top', 'single', 8), ('w:left', 'nil', 0), ('w:bottom', 'single', 8),
                         ('w:right', 'nil', 0), ('w:insideH', 'nil', 0), ('w:insideV', 'nil', 0)):
        b.append(border(tag, val, sz))
    # schema order inside tblPr: tblStyle, ..., tblW, jc, tblCellSpacing, tblInd, tblBorders, shd, tblLayout,
    # tblCellMar, tblLook
    ind = tblpr.find(qn('w:tblInd'))
    if ind is not None:
        ind.addnext(b)
    else:
        tblpr.insert(0, b)
    # header-row rule and table text size via conditional formatting
    for sp in el.findall(qn('w:tblStylePr')):
        el.remove(sp)
    sp = OxmlElement('w:tblStylePr')
    sp.set(qn('w:type'), 'firstRow')
    tcpr = OxmlElement('w:tcPr')
    tcb = OxmlElement('w:tcBorders')
    tcb.append(border('w:top', 'single', 8))
    tcb.append(border('w:bottom', 'single', 4))
    tcpr.append(tcb)
    va = OxmlElement('w:vAlign')
    va.set(qn('w:val'), 'bottom')
    tcpr.append(va)
    sp.append(tcpr)
    el.append(sp)
    rpr = el.find(qn('w:rPr'))
    if rpr is None:
        rpr = OxmlElement('w:rPr')
        el.insert(list(el).index(tblpr), rpr)
    set_fonts(rpr, 10)


def page_number_footer(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), ' PAGE ')
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = '1'
    r.append(t)
    fld.append(r)
    p._p.append(fld)
    set_fonts(r.get_or_add_rPr() if hasattr(r, 'get_or_add_rPr') else _rpr(r), 11)


def _rpr(r):
    rpr = r.find(qn('w:rPr'))
    if rpr is None:
        rpr = OxmlElement('w:rPr')
        r.insert(0, rpr)
    return rpr


def line_numbering(sectpr):
    for old in sectpr.findall(qn('w:lnNumType')):
        sectpr.remove(old)
    ln = OxmlElement('w:lnNumType')
    ln.set(qn('w:countBy'), '1')
    ln.set(qn('w:distance'), '284')        # 5 mm from the text
    ln.set(qn('w:restart'), 'continuous')
    # CT_SectPr order: ... pgSz, pgMar, paperSrc, pgBorders, lnNumType, pgNumType, cols, ...
    anchor = None
    for tag in ('w:pgBorders', 'w:paperSrc', 'w:pgMar', 'w:pgSz'):
        anchor = sectpr.find(qn(tag))
        if anchor is not None:
            break
    if anchor is not None:
        anchor.addnext(ln)
    else:
        sectpr.append(ln)


def build(out: Path, spacing: float = 1.5):
    with tempfile.TemporaryDirectory() as td:
        base = Path(td) / 'default.docx'
        with open(base, 'wb') as fh:
            subprocess.run([str(PANDOC), '--print-default-data-file', 'reference.docx'], stdout=fh, check=True)
        doc = Document(str(base))

    # ---- document defaults: Times New Roman 12 pt, en-GB
    styles_el = doc.styles.element
    rpr_default = styles_el.find(qn('w:docDefaults')).find(qn('w:rPrDefault')).find(qn('w:rPr'))
    set_fonts(rpr_default, 12)
    lang = rpr_default.find(qn('w:lang'))
    if lang is not None:
        lang.set(qn('w:val'), 'en-GB')
    ppr_default = styles_el.find(qn('w:docDefaults')).find(qn('w:pPrDefault')).find(qn('w:pPr'))
    for sp in ppr_default.findall(qn('w:spacing')):
        ppr_default.remove(sp)
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:after'), '0')
    ppr_default.append(sp)

    L = WD_ALIGN_PARAGRAPH.LEFT
    para_style(doc, 'Normal', size=12, align=L, before=0, after=0, line=spacing)
    for name in ('Body Text', 'First Paragraph'):
        para_style(doc, name, size=12, align=L, before=0, after=6, line=spacing)
    para_style(doc, 'Compact', size=12, align=L, before=0, after=0, line=spacing)
    para_style(doc, 'Title', size=16, bold=True, align=L, before=0, after=12, line=spacing)
    para_style(doc, 'Subtitle', size=13, bold=False, align=L, before=0, after=6, line=spacing)
    doc.styles['Subtitle'].element.find(qn('w:pPr')).remove(
        doc.styles['Subtitle'].element.find(qn('w:pPr')).find(qn('w:numPr'))) \
        if doc.styles['Subtitle'].element.find(qn('w:pPr')).find(qn('w:numPr')) is not None else None
    para_style(doc, 'Author', size=12, bold=False, italic=False, align=L, before=0, after=12, line=spacing)
    para_style(doc, 'Date', size=12, align=L, before=0, after=6, line=spacing)
    para_style(doc, 'Abstract Title', size=12, bold=True, align=L, before=12, after=6, line=spacing)
    para_style(doc, 'Abstract', size=12, align=L, before=0, after=6, line=spacing)
    para_style(doc, 'Keywords', size=12, align=L, before=6, after=12, line=spacing, base='Body Text')
    for lvl, (size, italic, before) in enumerate(((12, False, 18), (12, False, 12), (12, True, 12),
                                                  (12, True, 6)), start=1):
        st = para_style(doc, f'Heading {lvl}', size=size, bold=True, italic=italic, align=L, before=before,
                        after=6, line=spacing, keep_next=True)
        try:
            ch = doc.styles[f'Heading {lvl} Char']
            set_fonts(ch.element, size)
            ch.font.bold = True
            ch.font.italic = italic
            strip_colour(ch)
        except KeyError:
            pass
    for name in ('Title Char', 'Subtitle Char'):
        try:
            set_fonts(doc.styles[name].element)
            strip_colour(doc.styles[name])
        except KeyError:
            pass
    # captions: 10 pt, upright, left-aligned; table captions sit above the table and keep with it
    para_style(doc, 'Caption', size=10, italic=False, align=L, before=6, after=12, line=spacing)
    para_style(doc, 'Image Caption', size=10, italic=False, align=L, before=6, after=12, line=spacing)
    para_style(doc, 'Table Caption', size=10, italic=False, align=L, before=12, after=6, line=spacing,
               keep_next=True)
    C = WD_ALIGN_PARAGRAPH.CENTER
    para_style(doc, 'Figure', align=C, before=12, after=0, line=1.0)
    para_style(doc, 'Captioned Figure', align=C, before=12, after=0, line=1.0, keep_next=True)
    para_style(doc, 'Table Text', size=10, align=L, before=1, after=1, line=1.0, base='Normal')
    para_style(doc, 'Equation', align=C, before=3, after=3, line=1.0, base='Normal')
    para_style(doc, 'Equation Number', align=WD_ALIGN_PARAGRAPH.RIGHT, before=3, after=3, line=1.0, base='Normal')
    para_style(doc, 'Bibliography', size=12, align=L, before=0, after=3, line=spacing)
    # "[12]<tab>Author, ..." with the text in a hanging column 1 cm from the margin
    bib = find_style(doc, 'Bibliography')
    bib.paragraph_format.left_indent = Mm(10)
    bib.paragraph_format.first_line_indent = Mm(-10)
    bib.paragraph_format.tab_stops.add_tab_stop(Mm(10))
    para_style(doc, 'Footnote Text', size=10, align=L, before=0, after=0, line=1.0)
    para_style(doc, 'Block Text', size=12, align=L, before=6, after=6, line=spacing)
    para_style(doc, 'Footer', size=11, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=0, line=1.0, base='Normal')
    try:
        hl = doc.styles['Hyperlink']
        hl.font.color.rgb = RGBColor(0x1F, 0x3A, 0x93)
        hl.font.underline = False
    except KeyError:
        pass
    three_rule_table_style(doc)

    # ---- page set-up: A4, 25 mm margins, footer page number, continuous line numbers
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(25)
    sec.top_margin = sec.bottom_margin = Mm(25)
    sec.header_distance = sec.footer_distance = Mm(12.5)
    page_number_footer(sec)
    line_numbering(sec._sectPr)

    # the body of a reference docx is ignored by pandoc; keep it short but leave the sectPr
    body = doc.element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)
    p = doc.add_paragraph('Reference document for pandoc (styles only).', style='Body Text')
    body.remove(p._p)
    body.insert(0, p._p)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = ''
    doc.core_properties.author = ''
    doc.save(str(out))
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--spacing', type=float, default=1.5, help='line spacing multiple (1.5; 2 for submission)')
    ap.add_argument('--out', type=Path, default=DEFAULT_OUT)
    a = ap.parse_args()
    print(build(a.out, a.spacing))


if __name__ == '__main__':
    sys.exit(main())
